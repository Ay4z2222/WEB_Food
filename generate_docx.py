import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_with_badge(doc, text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = 'Tahoma'
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(234, 88, 12) # Orange 600
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    return h

def add_callout(doc, text, title="คำแนะนำการใช้งาน:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "FFF7ED") # light orange
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border orange
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="EA580C"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"💡 {title} ")
    r_title.bold = True
    r_title.font.name = 'Tahoma'
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(194, 65, 12)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Tahoma'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(67, 20, 7)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def create_table_guide(doc, data):
    """Create styled table with headers and explanation columns."""
    table = doc.add_table(rows=len(data) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Column widths: Badge (0.8 in), Name (2.0 in), Description (3.7 in)
    widths = [Inches(0.8), Inches(2.0), Inches(3.7)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    # Header Row
    headers = ["หมายเลข", "ชื่อปุ่ม / ส่วนประกอบ", "หน้าที่และการทำงานของระบบ"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "18181B") # Dark Zinc
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = 'Tahoma'
            run.font.size = Pt(10)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, item in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "FAFAFA" if row_idx % 2 == 0 else "FFFFFF"

        for col_idx in range(3):
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Badge
        row_cells[0].text = item[0]
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.runs[0].font.name = 'Tahoma'
        p0.runs[0].font.size = Pt(11)
        p0.runs[0].bold = True
        p0.runs[0].font.color.rgb = RGBColor(234, 88, 12)

        # Name
        row_cells[1].text = item[1]
        p1 = row_cells[1].paragraphs[0]
        p1.runs[0].font.name = 'Tahoma'
        p1.runs[0].font.size = Pt(9.5)
        p1.runs[0].bold = True
        p1.runs[0].font.color.rgb = RGBColor(24, 24, 27)

        # Description
        row_cells[2].text = item[2]
        p2 = row_cells[2].paragraphs[0]
        p2.runs[0].font.name = 'Tahoma'
        p2.runs[0].font.size = Pt(9.5)
        p2.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    # Subtle border for all cells
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                f'<w:left w:val="none"/>'
                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def build_manual():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title Page / Header Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_pre = title_p.add_run("คู่มือการปฏิบัติงาน & เอกสารแนะนำระบบ")
    r_pre.font.name = 'Tahoma'
    r_pre.font.size = Pt(11)
    r_pre.font.color.rgb = RGBColor(234, 88, 12)
    r_pre.bold = True

    title_main = doc.add_paragraph()
    title_main.paragraph_format.space_before = Pt(0)
    title_main.paragraph_format.space_after = Pt(4)
    r_title = title_main.add_run("ระบบสั่งอาหารบุฟเฟต์หมูกระทะ (QR Buffet Ordering)")
    r_title.font.name = 'Tahoma'
    r_title.font.size = Pt(20)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    r_sub = sub_p.add_run("คู่มืออธิบายการใช้งานแบบละเอียด พร้อมภาพประกอบและลูกศรชี้ระบุหน้าที่ทุกส่วนประกอบ\nครอบคลุม 3 ส่วนงาน: พนักงานแคชเชียร์หน้าร้าน, ห้องครัวจัดจาน (KDS), และลูกค้าสแกนสั่งอาหาร")
    r_sub.font.name = 'Tahoma'
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ----------------------------------------------------
    # SECTION 1: CASHIER & TABLE MANAGEMENT
    # ----------------------------------------------------
    add_heading_with_badge(doc, "1. หน้าจอแคชเชียร์และผังโต๊ะ (Cashier & Tables POS)", level=1)
    
    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_after = Pt(8)
    r = desc_p.add_run("หน้าจอนี้เปิดใช้งานบนเครื่องคอมพิวเตอร์หรือแท็บเล็ตประจำเคาน์เตอร์แคชเชียร์ เพื่อดูแลภาพรวมของร้านทั้งหมด ดูสถานะโต๊ะที่กำลังรับประทาน จัดการสต็อกของหมด และเปิดโต๊ะให้ลูกค้าใหม่")
    r.font.name = 'Tahoma'
    r.font.size = Pt(10)

    # Image with annotations
    cashier_img = r"C:\web\mookata-buffet\annotated_cashier.png"
    if os.path.exists(cashier_img):
        doc.add_picture(cashier_img, width=Inches(6.4))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("รูปที่ 1: ภาพรวมหน้าจอผังโต๊ะแคชเชียร์ พร้อมหมายเลขชี้ระบุฟังก์ชันการทำงาน")
        r_cap.font.name = 'Tahoma'
        r_cap.font.size = Pt(8.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    # Table explanation
    cashier_data = [
        ("①", "แถบสลับมุมมอง (View Switcher)", "ใช้สำหรับผู้ดูแลระบบหรือแคชเชียร์ในการสลับดูหน้าจอลูกค้า, จอห้องครัว, และจอแคชเชียร์ได้อย่างสะดวกบนหน้าจอเดียว"),
        ("②", "สรุปจำนวนโต๊ะ (Table Counter)", "แสดงจำนวนโต๊ะว่าง (สีเขียว) และโต๊ะที่กำลังรับประทานอยู่ (สีส้ม) แบบ Real-time ช่วยให้จัดคิวลูกค้าหน้าร้านได้ทันที"),
        ("③", "ปุ่มจัดการของหมด (Stock Toggle)", "กดเพื่อเปิดแถบเมนูปิด-เปิดวัตถุดิบ เมื่อของหมดระบบจะสั่งปิดการขายไปยังมือถือของลูกค้าทุกเครื่องพร้อมกันทันที"),
        ("④", "การ์ดโต๊ะว่าง (Available Table)", "แสดงสถานะไฟสีเขียวว่าง พร้อมปุ่ม [เปิดโต๊ะ] กดเพื่อเลือกรอบบุฟเฟต์และจำนวนลูกค้า เพื่อสร้าง QR Code"),
        ("⑤", "การ์ดโต๊ะกำลังทาน (Occupied Table)", "แสดงจำนวนลูกค้า, เวลานับถอยหลังของรอบ 90 นาที (ตัวเลขจะเปลี่ยนเป็นสีแดงเตือนเมื่อใกล้หมดเวลา), ปุ่มดู [QR] และปุ่ม [ปิดโต๊ะ/เช็คบิล]"),
    ]
    create_table_guide(doc, cashier_data)

    add_callout(
        doc,
        "เมื่อลูกค้าแจ้งเช็คบิล ให้แคชเชียร์กดปุ่ม [ปิดโต๊ะ] ระบบจะยกเลิกรหัส Token ของรอบนั้นทันที ป้องกันไม่ให้ลูกค้าเดิมนำลิงก์ไปแอบกดสั่งอาหารเล่นจากที่บ้านได้อย่างสมบูรณ์",
        title="ข้อควรปฏิบัติสำหรับแคชเชียร์:"
    )

    # ----------------------------------------------------
    # SECTION 2: QR CODE MODAL
    # ----------------------------------------------------
    doc.add_page_break()
    add_heading_with_badge(doc, "2. หน้าต่างสร้าง QR Code ประจำโต๊ะ (Dynamic QR Token)", level=1)

    desc_p2 = doc.add_paragraph()
    desc_p2.paragraph_format.space_after = Pt(8)
    r2 = desc_p2.add_run("เมื่อแคชเชียร์กด 'เปิดโต๊ะ' และระบุจำนวนลูกค้า ระบบจะสร้างรหัสลับเฉพาะรอบ (Token) และสร้าง QR Code บุฟเฟต์ 90 นาทีขึ้นมาทันที")
    r2.font.name = 'Tahoma'
    r2.font.size = Pt(10)

    qr_img = r"C:\web\mookata-buffet\annotated_qr_modal.png"
    if os.path.exists(qr_img):
        doc.add_picture(qr_img, width=Inches(6.4))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("รูปที่ 2: หน้าต่าง QR Code สำหรับให้ลูกค้านำมือถือมาสแกนสั่งอาหาร")
        r_cap.font.name = 'Tahoma'
        r_cap.font.size = Pt(8.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    qr_data = [
        ("①", "ชื่อโต๊ะและรหัสลับ (Table & Token)", "ระบุหมายเลขโต๊ะ เช่น 'โต๊ะ 02' และรหัส Session Token 6 หลัก เช่น #U1JG3F เพื่อใช้ยืนยันความถูกต้องของรอบการทาน"),
        ("②", "คิวอาร์โค้ดสั่งอาหาร (QR Code)", "โค้ดภาพความละเอียดสูง ลูกค้าสามารถใช้กล้องมือถือทุกรุ่น หรือแอป LINE สแกนเพื่อเข้าสู่เมนูร้านได้ทันทีโดยไม่ต้องโหลดแอปพลิเคชัน"),
        ("③", "ปุ่มจำลองการสแกน (Simulate Scan)", "กดปุ่มนี้เพื่อสลับเข้าสู่หน้าจอลูกค้าของโต๊ะนี้ทันทีบนหน้าจอคอมพิวเตอร์ เหมาะสำหรับทดสอบหรือช่วยลูกค้ากดสั่งแทน"),
        ("④", "ปุ่มพิมพ์สลิป QR Code (Print Slip)", "กดเพื่อส่งคำสั่งพิมพ์ QR Code และเลขโต๊ะออกทางเครื่องพิมพ์สลิปความร้อน (Thermal Printer) สำหรับนำไปวางที่โต๊ะอาหาร"),
    ]
    create_table_guide(doc, qr_data)

    # ----------------------------------------------------
    # SECTION 3: CUSTOMER MOBILE VIEW
    # ----------------------------------------------------
    doc.add_page_break()
    add_heading_with_badge(doc, "3. หน้าจอสั่งอาหารของลูกค้า (Customer Mobile Web)", level=1)

    desc_p3 = doc.add_paragraph()
    desc_p3.paragraph_format.space_after = Pt(8)
    r3 = desc_p3.add_run("หน้าจอที่แสดงบนสมาร์ตโฟนของลูกค้าหลังจากสแกน QR Code ออกแบบให้ใช้งานด้วยมือเดียวได้สะดวก มีระบบนับเวลา 90 นาทีติดอยู่ด้านบน พร้อมระบบเรียกบริการด่วน")
    r3.font.name = 'Tahoma'
    r3.font.size = Pt(10)

    cust_img = r"C:\web\mookata-buffet\annotated_customer.png"
    if os.path.exists(cust_img):
        doc.add_picture(cust_img, width=Inches(3.3))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("รูปที่ 3: หน้าจอมือถือของลูกค้า พร้อมลูกศรชี้องค์ประกอบสำคัญ")
        r_cap.font.name = 'Tahoma'
        r_cap.font.size = Pt(8.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    cust_data = [
        ("①", "ข้อมูลโต๊ะและรหัสรอบ", "แสดงหมายเลขโต๊ะที่ลูกค้ากำลังนั่งอยู่ และป้ายยืนยันว่าเป็นบุฟเฟต์หมูกระทะ"),
        ("②", "นาฬิกานับถอยหลัง 90 นาที (Server Timer)", "ยึดเวลาจากเซิร์ฟเวอร์หลัก (ลูกค้าแอบปรับเวลาในมือถือไม่ได้) เหลือ 15 นาทีจะกะพริบเตือนสีส้ม และล็อคปุ่มสั่งอัตโนมัติเมื่อครบ 90 นาที"),
        ("③", "ป้ายแจ้งเตือนของหวานตักเอง", "แจ้งให้ลูกค้าทราบว่าไอศกรีม 6 รส, รวมมิตรน้ำกะทิ และผลไม้ บริการตัวเองได้ที่บาร์ของหวานด้านหลังร้าน โดยไม่ต้องสั่งผ่านแอป"),
        ("④", "แถบเลือกหมวดหมู่อาหาร", "แบ่งหมวดหมู่ชัดเจน: หมู/เนื้อ, ผัก/เห็ด, ของทานเล่น, เครื่องดื่ม, ของหวาน และบริการเสริม"),
        ("⑤", "การ์ดอาหาร & ปุ่มเพิ่มถาด", "แสดงรูปอาหารจริง, คำอธิบาย, ป้ายบุฟเฟต์ 0.- และปุ่ม [+] / [-] ปรับจำนวนถาดตามต้องการ"),
        ("⑥", "ปุ่มบริการด่วน (Quick Service Bell)", "กดเพื่อเรียกพนักงานด่วน: ขอเติมถ่าน, ขอเปลี่ยนกระทะ, ขอเติมน้ำซุป หรือเรียกพนักงานที่โต๊ะ"),
        ("⑦", "ปุ่มประวัติการสั่ง (Order History)", "ลูกค้าสามารถกดดูได้ว่าตนเองสั่งอะไรไปแล้วบ้าง และแต่ละจานเสิร์ฟถึงโต๊ะแล้วหรือยัง"),
        ("⑧", "ปุ่มตะกร้า & ส่งออเดอร์เข้าครัว", "แสดงจำนวนถาดทั้งหมดที่เลือกไว้ เมื่อกดแล้วจะมีหน้าสรุปยืนยัน และกดส่งเข้าครัวพร้อมเอฟเฟกต์พลุฉลอง"),
    ]
    create_table_guide(doc, cust_data)

    add_callout(
        doc,
        "หากวัตถุดิบรายการใดหมด การ์ดอาหารนั้นจะกลายเป็นสีเทาและแสดงป้าย 'ของหมด' อัตโนมัติ โดยที่ลูกค้าจะไม่สามารถกดปุ่มบวกเพิ่มลงตะกร้าได้ เพื่อป้องกันปัญหาการสั่งอาหารที่ไม่มีของ",
        title="กลไกป้องกันของหมด (Out of Stock Control):"
    )

    # ----------------------------------------------------
    # SECTION 4: KITCHEN DISPLAY SYSTEM (KDS)
    # ----------------------------------------------------
    doc.add_page_break()
    add_heading_with_badge(doc, "4. หน้าจอห้องครัวและแผนกเตรียมของ (Kitchen Display - KDS)", level=1)

    desc_p4 = doc.add_paragraph()
    desc_p4.paragraph_format.space_after = Pt(8)
    r4 = desc_p4.add_run("หน้าจอนี้เปิดบนแท็บเล็ตหรือจอโทรทัศน์ติดผนังในห้องครัว แผนกสไลด์เนื้อ และจุดทอดของทานเล่น เพื่อรับออเดอร์ใหม่แบบสดๆ โดยไม่ต้องใช้กระดาษจด")
    r4.font.name = 'Tahoma'
    r4.font.size = Pt(10)

    kit_img = r"C:\web\mookata-buffet\annotated_kitchen.png"
    if os.path.exists(kit_img):
        doc.add_picture(kit_img, width=Inches(6.4))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run("รูปที่ 4: หน้าจอรับออเดอร์ห้องครัวแบบ Real-time พร้อมแถบเตือนบริการด่วน")
        r_cap.font.name = 'Tahoma'
        r_cap.font.size = Pt(8.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    kit_data = [
        ("①", "หัวข้อจอครัว & ไฟสถานะ", "แสดงชื่อระบบ KDS และมีจุดไฟสีเขียวกะพริบแสดงว่าระบบกำลังเชื่อมต่อสัญญาณ Real-time กับหน้าร้านอยู่ตลอดเวลา"),
        ("②", "ปุ่มเปิด/ปิดเสียงกระดิ่งเตือน", "เมื่อมีออเดอร์ใหม่หรือมีลูกค้าเรียกถ่าน ระบบจะส่งเสียงกระดิ่ง 'ติ๊งหน่อง' สามารถกดเปิดหรือปิดเสียงได้ที่ปุ่มนี้"),
        ("③", "แถบเตือนบริการด่วน (Urgent Service)", "แสดงคำขอเร่งด่วน เช่น โต๊ะ 02 ขอเติมถ่าน หรือขอเปลี่ยนกระทะ เมื่อพนักงานนำไปทำให้ลูกค้าแล้วให้กดปุ่ม [✓ จัดการแล้ว] แถบจะหายไป"),
        ("④", "การ์ดออเดอร์อาหาร (Order Ticket)", "ระบุหมายเลขโต๊ะชัดเจน, เวลาที่ลูกค้ารอ (นาที), รายการอาหารและจำนวนถาดที่ต้องจัด หากรอนานเกิน 5 นาทีการ์ดจะเปลี่ยนเป็นขอบสีแดงเตือน"),
        ("⑤", "ปุ่มปรับสถานะการจัดส่ง", "กดปุ่ม [กำลังจัดของ] เพื่อแจ้งลูกค้าว่าครัวกำลังทำ และกดปุ่ม [เสิร์ฟแล้ว] เมื่อเด็กเสิร์ฟนำถาดอาหารไปวางที่โต๊ะลูกค้าเรียบร้อยแล้ว"),
    ]
    create_table_guide(doc, kit_data)

    # ----------------------------------------------------
    # SECTION 5: TROUBLESHOOTING & BEST PRACTICES
    # ----------------------------------------------------
    doc.add_page_break()
    add_heading_with_badge(doc, "5. แนวทางการปฏิบัติงานและแก้ปัญหาเบื้องต้น (Best Practices)", level=1)

    faq_p = doc.add_paragraph()
    faq_p.paragraph_format.space_after = Pt(6)
    r_f1 = faq_p.add_run("1. เมื่อลูกค้าแจ้งว่าสแกน QR Code แล้วหน้าจอขึ้นว่า 'หมดอายุ':\n")
    r_f1.bold = True
    r_f1.font.name = 'Tahoma'
    r_f1.font.size = Pt(10)
    r_f1.font.color.rgb = RGBColor(15, 23, 42)
    r_f2 = faq_p.add_run("สาเหตุ: โต๊ะนั้นอาจจะยังไม่ได้กด 'เปิดโต๊ะ' ในระบบแคชเชียร์ หรือเป็น QR Code ของลูกค้ารอบก่อนหน้าที่เช็คบิลไปแล้ว\nวิธีแก้: ให้แคชเชียร์ไปที่หน้าจอ POS แล้วกดเปิดโต๊ะใหม่ จากนั้นส่ง QR Code ใบใหม่ให้ลูกค้าสแกน\n\n")
    r_f2.font.name = 'Tahoma'
    r_f2.font.size = Pt(9.5)
    r_f2.font.color.rgb = RGBColor(71, 85, 105)

    r_f3 = faq_p.add_run("2. เมื่ออาหารในครัวหมดกะทันหันระหว่างวัน:\n")
    r_f3.bold = True
    r_f3.font.name = 'Tahoma'
    r_f3.font.size = Pt(10)
    r_f3.font.color.rgb = RGBColor(15, 23, 42)
    r_f4 = faq_p.add_run("วิธีแก้: พนักงานแคชเชียร์สามารถกดปุ่ม [จัดการของหมด] ที่มุมขวาบนของหน้าจอแคชเชียร์ แล้วกดปิดรายการอาหารนั้นทันที มือถือของลูกค้าทุกโต๊ะจะถูกล็อคไม่ให้สั่งเมนูนั้นโดยอัตโนมัติ\n\n")
    r_f4.font.name = 'Tahoma'
    r_f4.font.size = Pt(9.5)
    r_f4.font.color.rgb = RGBColor(71, 85, 105)

    r_f5 = faq_p.add_run("3. การแจ้งเตือนเมื่อใกล้หมดเวลา 90 นาที:\n")
    r_f5.bold = True
    r_f5.font.name = 'Tahoma'
    r_f5.font.size = Pt(10)
    r_f5.font.color.rgb = RGBColor(15, 23, 42)
    r_f6 = faq_p.add_run("เมื่อเวลาของโต๊ะเหลือต่ำกว่า 15 นาที แถบเวลาบนจอลูกค้าจะเปลี่ยนเป็นสีส้ม/แดงกะพริบเตือนให้สั่งรอบสุดท้าย และเมื่อเวลาหมดลง (00:00) ปุ่มส่งออเดอร์จะถูกปิดการทำงาน พนักงานหน้าร้านสามารถนำบิลไปวางเช็คบิลที่โต๊ะได้ทันที\n")
    r_f6.font.name = 'Tahoma'
    r_f6.font.size = Pt(9.5)
    r_f6.font.color.rgb = RGBColor(71, 85, 105)

    out_file1 = r"C:\web\mookata-buffet\Mookata_Buffet_User_Manual.docx"
    doc.save(out_file1)
    print(f"Manual successfully saved to: {out_file1}")

if __name__ == '__main__':
    build_manual()
